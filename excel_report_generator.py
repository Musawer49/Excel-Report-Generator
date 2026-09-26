from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
import pandas as pd

data = pd.read_excel("sales_data.xlsx")

data["Total"] = data["Quantity"] * data["Price"]

total_sales = data["Total"].sum()
total_quantity_sales = data["Quantity"].sum()

category_sales = data.groupby("Category")["Total"].sum()
category_report = category_sales.reset_index(name="Total Sales")

product_sales = data.groupby("Product")["Total"].sum()

product_report = product_sales.reset_index(name="Total Sales")
product_report = product_report.sort_values("Total Sales", ascending=False)

product_sales = data.groupby("Product")["Total"].sum()
top_product = product_sales.idxmax()

report = pd.DataFrame(
    [
        ["Total Quantity Sold", total_quantity_sales],
        ["Total Sales", total_sales],
        ["Top Product", top_product],
    ],
    columns=["Sales", "Values"]
)

with pd.ExcelWriter(
    "sales_report.xlsx",
    engine="openpyxl"
) as writer:

    report.to_excel(writer, sheet_name="Summary", index=False)
    category_report.to_excel(writer, sheet_name="Category Sales", index=False)
    data.to_excel(writer, sheet_name="Raw Data", index=False)

    product_report.to_excel(
        writer,
        sheet_name="Product Sales",
        index=False
    )

    workbook = writer.book

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    summary = writer.sheets["Summary"]

    summary.column_dimensions["A"].width = 25
    summary.column_dimensions["B"].width = 20

    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for row in summary.iter_rows():
        for cell in row:
            cell.border = thin_border

    summary["B2"].number_format = "#,##0"
    summary["B3"].number_format = "#,##0"

    category = writer.sheets["Category Sales"]

    category.column_dimensions["A"].width = 20
    category.column_dimensions["B"].width = 20

    for cell in category[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for row in category.iter_rows():
        for cell in row:
            cell.border = thin_border

    for row in category.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "#,##0"

    chart = BarChart()
    chart.title = "Sales by Category"
    chart.y_axis.title = "Total Sales"
    chart.x_axis.title = "Category"

    data_ref = Reference(
        category,
        min_col=2,
        min_row=1,
        max_row=category.max_row
    )

    categories_ref = Reference(
        category,
        min_col=1,
        min_row=2,
        max_row=category.max_row
    )

    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(categories_ref)

    category.add_chart(chart, "D2")

    raw = writer.sheets["Raw Data"]

    for cell in raw[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for row in raw.iter_rows():
        for cell in row:
            cell.border = thin_border

    raw.freeze_panes = "A2"

    raw.column_dimensions["A"].width = 15
    raw.column_dimensions["B"].width = 18
    raw.column_dimensions["C"].width = 18
    raw.column_dimensions["D"].width = 12
    raw.column_dimensions["E"].width = 15
    raw.column_dimensions["F"].width = 15

    product = writer.sheets["Product Sales"]

    product.column_dimensions["A"].width = 20
    product.column_dimensions["B"].width = 20

    for cell in product[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for row in product.iter_rows():
        for cell in row:
            cell.border = thin_border

    for row in product.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "#,##0"

    product_chart = BarChart()
    product_chart.title = "Sales by Product"
    product_chart.y_axis.title = "Total Sales"
    product_chart.x_axis.title = "Product"

    product_data = Reference(
        product,
        min_col=2,
        min_row=1,
        max_row=product.max_row
    )

    product_categories = Reference(
        product,
        min_col=1,
        min_row=2,
        max_row=product.max_row
    )

    product_chart.add_data(product_data, titles_from_data=True)
    product_chart.set_categories(product_categories)

    product.add_chart(product_chart, "D2")
